# -*- coding: utf-8 -*-
"""Tests del kit. Ejecutar desde la raíz:  .venv/bin/python -m unittest -v

No tocan el .env ni materias/ reales: los de extremo a extremo copian
_plantillas/ a un directorio temporal con su propio .env.
"""
import os
import shutil
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "_plantillas"))

ENV = textwrap.dedent("""\
    NOMBRE_ESTUDIANTE="Ana Pérez"
    INSTITUCION="Universidad Ejemplo"
    CARRERA="Carrera de Software"
    SUFIJO_ARCHIVOS="Perez_Ana"
    """)

import apa_docx  # noqa: E402  (importar no debe exigir el .env)
import comprobar as comp  # noqa: E402
import esqueleto as esq  # noqa: E402
import ficha  # noqa: E402
import rubrica as rub  # noqa: E402


class TestRubrica(unittest.TestCase):
    def test_numero(self):
        self.assertEqual(rub.numero("2,5 pts"), 2.5)
        self.assertEqual(rub.numero("20 %"), 20.0)
        self.assertIsNone(rub.numero("<!-- -->"))

    def test_leer_tabla_antigua_sin_id(self):
        texto = textwrap.dedent("""\
            # A1 — X

            ## Rúbrica

            | Criterio | Puntaje | Dónde se cumple |
            |---|---|---|
            | Diseño del modelo | 3 | |
            | Ortografía | 1 | |

            ## Qué se entrega
            """)
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "TAREA.md"
            ruta.write_text(texto, encoding="utf-8")
            criterios = rub.leer(ruta)
            self.assertEqual([c.id for c in criterios], ["C1", "C2"])
            self.assertEqual(criterios[0].puntaje, 3)
            self.assertTrue(criterios[1].transversal)
            rub.escribir(ruta, criterios)
            nuevo = ruta.read_text(encoding="utf-8")
            self.assertIn("| ID | Criterio | Puntaje | Dónde se cumple |", nuevo)
            self.assertIn("## Qué se entrega", nuevo)
            self.assertEqual(len(rub.leer(ruta)), 2)

    def test_referencias_no_es_transversal(self):
        c = rub.Criterio("C4", "Referencias en formato APA")
        self.assertTrue(c.de_referencias)
        self.assertFalse(c.transversal)

    def test_importar_matriz_docx(self):
        from docx import Document
        doc = Document()
        t = doc.add_table(rows=1, cols=3)
        for celda, texto in zip(t.rows[0].cells, ["Criterio", "Excelente (4)", "Regular (2)"]):
            celda.text = texto
        for fila in (["Explica el concepto", "a", "b"], ["Implementa", "c", "d"], ["Total", "", ""]):
            for celda, texto in zip(t.add_row().cells, fila):
                celda.text = texto
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "r.docx"
            doc.save(ruta)
            criterios = rub.importar_docx(ruta)
        self.assertEqual([(c.id, c.texto, c.puntaje) for c in criterios],
                         [("C1", "Explica el concepto", 4), ("C2", "Implementa", 4)])


class TestFicha(unittest.TestCase):
    def test_fecha_y_extension(self):
        self.assertEqual(str(ficha.leer_fecha("2026-10-25")), "2026-10-25")
        self.assertEqual(str(ficha.leer_fecha("25/10/2026")), "2026-10-25")
        self.assertIsNone(ficha.leer_fecha("<!-- AAAA-MM-DD -->"))
        self.assertEqual(ficha.leer_extension("1.500 palabras"), 1500)
        self.assertEqual(ficha.leer_extension("4-6 páginas"), 5 * ficha.PALABRAS_POR_PAGINA)

    def test_celdas_con_barra_escapada(self):
        self.assertEqual(ficha.celdas("| a | b \\| c |"), ["a", "b | c"])


class TestBibliografia(unittest.TestCase):
    def setUp(self):
        self.bib = apa_docx.Bibliografia()
        self.bib.agregar("ramos", ["Ramos, J.", "Pérez, A."], 2023, "Bases de datos", "Alfaomega")
        self.bib.agregar("lee", ["Lee, K.", "Kim, S.", "Park, J."], 2021, "Indexing",
                         "Journal of Databases", detalles="12(3), 45–60", tipo="articulo")

    def test_citas(self):
        self.assertEqual(self.bib.cita("ramos"), "(Ramos y Pérez, 2023)")
        self.assertEqual(self.bib.cita("ramos", "p. 15"), "(Ramos y Pérez, 2023, p. 15)")
        self.assertEqual(self.bib.cita("ramos", "lee"), "(Lee et al., 2021; Ramos y Pérez, 2023)")
        self.assertEqual(self.bib.cita_narrativa("lee"), "Lee et al. (2021)")
        with self.assertRaises(KeyError):
            self.bib.cita("no-existe")

    def test_entradas(self):
        self.assertEqual(self.bib.entrada("ramos"),
                         "Ramos, J. y Pérez, A. (2023). *Bases de datos*. Alfaomega.")
        self.assertEqual(self.bib.entrada("lee"), "Lee, K., Kim, S. y Park, J. (2021). Indexing. "
                                                  "*Journal of Databases*, 12(3), 45–60.")

    def test_solo_citadas_y_en_orden(self):
        self.bib.cita("ramos")
        self.assertEqual(self.bib.ordenadas(), ["ramos"])
        self.assertEqual(self.bib.no_citadas(), ["lee"])


class TestCitasEnTexto(unittest.TestCase):
    def test_detecta_citas(self):
        texto = ("Según Ramos y Pérez (2023) y Lee et al. (2021), algo (Oracle, 2024; "
                 "Ministerio de Educación, s.f., p. 3). Ver Figura 1 (captura). ISO (9001).")
        self.assertEqual(sorted(set(comp._citas(texto))),
                         [("Lee", "2021"), ("Ministerio de Educación", "s.f."),
                          ("Oracle", "2024"), ("Ramos", "2023")])

    def test_coincide(self):
        self.assertTrue(comp._coincide("Ramos", "2023", "Ramos, J. y Pérez, A. (2023). *X*."))
        self.assertFalse(comp._coincide("Ramos", "2022", "Ramos, J. (2023). X."))


class TestEsqueleto(unittest.TestCase):
    def test_bloque(self):
        criterios = [rub.Criterio("C1", "Explica el concepto", 6),
                     rub.Criterio("C2", "Implementa el código", 2),
                     rub.Criterio("C3", "Ortografía y redacción", 1),
                     rub.Criterio("C4", "Referencias APA", 1)]
        bloque, refs = esq.bloque(criterios, 1000)
        self.assertIn("criterios='C1'", bloque)
        self.assertIn("~600 palabras", bloque)  # 800 × 6/8
        self.assertIn("~200 palabras", bloque)  # 800 × 2/8
        self.assertIn("transversal(doc, 'C3')", bloque)
        self.assertEqual(refs, "C4")

    def test_no_pisa_texto_escrito(self):
        self.assertFalse(esq._tiene_texto('# x\ntitulo(doc, "A", 1)\nparrafo(doc, "")\n'))
        self.assertTrue(esq._tiene_texto('parrafo(doc, "Ya escrito")\n'))


class TestExtremoAExtremo(unittest.TestCase):
    """Crea materia y tarea, importa la rúbrica, genera y comprueba."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.kit = Path(self.tmp.name)
        shutil.copytree(RAIZ / "_plantillas", self.kit / "_plantillas",
                        ignore=shutil.ignore_patterns("__pycache__"))
        (self.kit / ".env").write_text(ENV, encoding="utf-8")
        self.entorno = {k: v for k, v in os.environ.items() if k != "RUBRIC_DRIVEN_ENV"}

    def tearDown(self):
        self.tmp.cleanup()

    def tareas(self, *args, ok=True):
        r = subprocess.run([sys.executable, "_plantillas/tareas.py", *args], cwd=self.kit,
                           capture_output=True, text=True, env=self.entorno)
        if ok and r.returncode != 0:
            self.fail(f"tareas.py {' '.join(args)} → {r.returncode}\n{r.stdout}\n{r.stderr}")
        return r

    def test_flujo_completo(self):
        self.tareas("materia", "sgbd1", "Sistemas de Bases de Datos I")
        readme = self.kit / "materias/sgbd1/README.md"
        readme.write_text(readme.read_text(encoding="utf-8").replace(
            "| **Docente** | <!-- nombre --> |", "| **Docente** | Ing. Luis Gómez |"), encoding="utf-8")
        self.tareas("tarea", "sgbd1", "A15", "Procedimientos almacenados")
        carpeta = self.kit / "materias/sgbd1/A15_Procedimientos_Almacenados"
        tarea_md = carpeta / "TAREA.md"
        tarea_md.write_text(tarea_md.read_text(encoding="utf-8")
                            .replace("<!-- AAAA-MM-DD -->", "2026-10-25")
                            .replace("| ID | Criterio | Puntaje | Dónde se cumple |\n|---|---|---|---|\n",
                                     "| ID | Criterio | Puntaje | Dónde se cumple |\n|---|---|---|---|\n"
                                     "| C1 | Explica el concepto | 3 | |\n"
                                     "| C2 | Implementa el código | 1 | |\n"
                                     "| C3 | Ortografía y redacción | 1 | |\n"
                                     "| C4 | Referencias APA | 1 | |\n"), encoding="utf-8")
        ruta = "materias/sgbd1/A15_Procedimientos_Almacenados"
        self.tareas("esqueleto", ruta)

        # Sin contenido: secciones vacías → la comprobación falla.
        r = self.tareas("generar", ruta, "--sin-pdf", ok=False)
        self.assertEqual(r.returncode, 1, r.stdout)
        self.assertIn("está vacía", r.stdout)

        script = carpeta / "generar_documento.py"
        texto = script.read_text(encoding="utf-8")
        texto = texto.replace(
            '# bib.agregar("clave", ["Apellido, A."], 2026, "Título de la obra", "Editorial")',
            'bib.agregar("ramos", ["Ramos, J."], 2023, "Bases de datos", "Alfaomega")')
        texto = texto.replace("criterios='C1')\nparrafo(doc, \"\")",
                              "criterios='C1')\nparrafo(doc, f\"Concepto {bib.cita('ramos')}. \" * 60)")
        texto = texto.replace("criterios='C2')\nparrafo(doc, \"\")",
                              "criterios='C2')\nparrafo(doc, \"Código del ejercicio. \" * 30)")
        texto = texto.replace('titulo(doc, "Conclusiones", 1)\nparrafo(doc, "")',
                              'titulo(doc, "Conclusiones", 1)\nparrafo(doc, "Fin. " * 20)')
        script.write_text(texto, encoding="utf-8")

        # El esqueleto ya no se reescribe sin --forzar.
        self.assertNotEqual(self.tareas("esqueleto", ruta, ok=False).returncode, 0)

        r = self.tareas("generar", ruta, "--sin-pdf", ok=False)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        tarea = tarea_md.read_text(encoding="utf-8")
        self.assertIn("| C1 | Explica el concepto | 3 | Explica el concepto (", tarea)
        self.assertIn("| C3 | Ortografía y redacción | 1 | todo el documento |", tarea)
        self.assertIn("| C4 | Referencias APA | 1 | Referencias |", tarea)

        docx = carpeta / "A15_Procedimientos_Almacenados_Perez_Ana.docx"
        from docx import Document
        doc = Document(str(docx))
        textos = [p.text for p in doc.paragraphs]
        self.assertIn("Ing. Luis Gómez", textos)            # portada con docente
        self.assertIn("25 de octubre de 2026", textos)       # y fecha de entrega
        self.assertIn("Ramos, J. (2023). Bases de datos. Alfaomega.", textos)
        estilos = {p.style.name for p in doc.paragraphs}
        self.assertIn("Heading 1", estilos)

        self.tareas("marcar", ruta, "lista")
        self.assertIn("| lista |", (self.kit / "materias/sgbd1/README.md").read_text(encoding="utf-8"))
        self.assertIn("A15_Procedimientos_Almacenados", self.tareas("estado").stdout)

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "sin LibreOffice")
    def test_pdf(self):
        self.tareas("materia", "m", "Materia")
        self.tareas("tarea", "m", "A1", "Prueba")
        ruta = "materias/m/A1_Prueba"
        r = self.tareas("generar", ruta, ok=False)
        self.assertIn("PDF generado", r.stdout, r.stdout + r.stderr)
        self.assertTrue((self.kit / ruta / "A1_Prueba_Perez_Ana.pdf").exists())


class TestImportarSinEnv(unittest.TestCase):
    def test_importar_no_lee_env(self):
        entorno = dict(os.environ, RUBRIC_DRIVEN_ENV="/no/existe/.env")
        r = subprocess.run([sys.executable, "-c", "import apa_docx; print('ok')"],
                           cwd=RAIZ / "_plantillas", capture_output=True, text=True, env=entorno)
        self.assertEqual(r.stdout.strip(), "ok", r.stderr)


if __name__ == "__main__":
    unittest.main()
